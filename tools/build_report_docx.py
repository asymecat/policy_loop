#!/usr/bin/env python3
"""从 docs/report.md 生成 docs/report.docx。

单一事实来源是 docs/report.md。**不要手工编辑 report.docx**——它是派生物，
随时可以被重新生成，任何手工修改都会在下次生成时丢失。

处理步骤：
  1. 读 docs/report.md，去掉手写的「## 目录」节（docx 用 Word 自动目录域，
     两份并存会变成两个目录）；
  2. 交给 pandoc 转 docx（--toc 生成的是真正的 TOC 域，Word 里可更新出页码）；
  3. 用 python-docx 后处理：把目录挪到封面页之后并另起一页、目录标题改中文、
     正文章节另起一页、设定中西文字体、表头跨页重复且单行不跨页拆分。

用法：
    python3 tools/build_report_docx.py                # 生成 docs/report.docx
    python3 tools/build_report_docx.py --out /tmp/x.docx
    python3 tools/build_report_docx.py --check        # 只比对，不写盘

依赖：pandoc（命令行）、python-docx。
"""

from __future__ import annotations

import argparse
import hashlib
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
SRC = DOCS / "report.md"
DST = DOCS / "report.docx"

# 样式 -> (西文字体, 中文字体)
FONT_SPEC = {
    "Normal": ("Times New Roman", "宋体"),
    "Body Text": ("Times New Roman", "宋体"),
    "First Paragraph": ("Times New Roman", "宋体"),
    "Compact": ("Times New Roman", "宋体"),
    "Block Text": ("Times New Roman", "宋体"),
    "Source Code": ("Consolas", "宋体"),
    "Title": ("Times New Roman", "黑体"),
    "Heading 1": ("Times New Roman", "黑体"),
    "Heading 2": ("Times New Roman", "黑体"),
    "Heading 3": ("Times New Roman", "黑体"),
}


def die(msg: str) -> None:
    print(f"错误：{msg}", file=sys.stderr)
    raise SystemExit(1)


def strip_manual_toc(text: str) -> str:
    """去掉 md 里手写的「## 目录」节。

    那份目录对在 GitHub 上读 md 是有用的，但 docx 走 pandoc 的自动目录域，
    两份并存会出现两个目录。只影响生成时的输入，不改动 report.md 本身。
    """
    head = re.search(r"^## 目录\s*$", text, re.M)
    if not head:
        return text
    rest = text[head.end():]
    tail = re.search(r"^---\s*$", rest, re.M)
    if not tail:
        return text
    return text[: head.start()] + rest[tail.end():].lstrip("\n")


def run_pandoc(build_md: Path, out_path: Path) -> None:
    if shutil.which("pandoc") is None:
        die("找不到 pandoc，请先安装（如 apt install pandoc）")
    subprocess.run(
        [
            "pandoc", build_md.name,
            "-o", str(out_path.resolve()),
            "--toc", "--toc-depth=3",
            "--resource-path", str(DOCS),
        ],
        check=True,
        cwd=DOCS,
    )


def set_fonts(doc) -> None:
    from docx.oxml.ns import qn

    for name, (latin, east_asian) in FONT_SPEC.items():
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        rfonts.set(qn("w:ascii"), latin)
        rfonts.set(qn("w:hAnsi"), latin)
        rfonts.set(qn("w:eastAsia"), east_asian)


def move_toc_and_paginate(doc) -> None:
    """把自动目录域从文档最前挪到封面页之后，并给目录/正文分页。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    body = doc.element.body
    sdt = body.find(qn("w:sdt"))
    if sdt is None:
        return

    # 目录标题：pandoc 写的是英文 "Table of Contents"
    for node in sdt.iter(qn("w:t")):
        if node.text and "Table of Contents" in node.text:
            node.text = "目录"

    # 封面页 = 标题 + 赛题行 + 信息表格 + 提交材料说明，目录接在它后面
    anchor = None
    for para in doc.paragraphs:
        if para.text.startswith("提交材料"):
            anchor = para._p
            break
    if anchor is None and doc.tables:
        anchor = doc.tables[0]._tbl

    body.remove(sdt)
    (anchor if anchor is not None else body).addnext(sdt)

    # 目录另起一页
    toc_para = sdt.find(".//" + qn("w:p"))
    if toc_para is not None:
        toc_para.get_or_add_pPr().get_or_add_pageBreakBefore()

    # 封面：标题与赛题行居中，标题放大
    filled = [p for p in doc.paragraphs if p.text.strip()]
    if filled:
        try:
            filled[0].style = doc.styles["Title"]  # 脱离 Heading 1，才不会进自动目录
        except KeyError:
            pass
        filled[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in filled[0].runs:
            run.font.size = Pt(22)
            run.font.bold = True
    if len(filled) > 1:
        filled[1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 正文第一章另起一页
    for para in doc.paragraphs:
        if para.style.name == "Heading 1" and para.text.strip().startswith("1 "):
            para.paragraph_format.page_break_before = True
            break


def extract_headings(md_text: str) -> list[tuple[int, str]]:
    """取出 1–3 级标题，用作目录域的预填内容。跳过代码块里的 `#` 注释。"""
    out: list[tuple[int, str]] = []
    in_code = False
    for line in md_text.splitlines():
        if line.lstrip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            continue
        m = re.match(r"^(#{1,3})\s+(\S.*?)\s*$", line)
        if m:
            # 目录条目里不该出现 markdown 记号
            out.append((len(m.group(1)), re.sub(r"[`*]", "", m.group(2)).strip()))
    return out[1:] if out else out  # 首条是封面标题，不进目录


def set_page_a4(doc) -> None:
    from docx.shared import Mm

    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(25.4)
        section.bottom_margin = Mm(25.4)
        section.left_margin = Mm(31.7)
        section.right_margin = Mm(25.4)


def fill_toc(doc, headings: list[tuple[int, str]]) -> None:
    """往 pandoc 的 TOC 域里填一份静态目录。

    pandoc 生成的域是空的（begin/instrText/separate/end 挤在同一个 run 里），
    不更新域就什么都看不到——LibreOffice 渲染出的 PDF 里目录页是空白的。
    这里把 separate 与 end 拆开、中间放上标题条目：任何查看器都能看到内容，
    而 Word 更新域时会把它们整段替换成带页码的真目录。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    sdt = doc.element.body.find(qn("w:sdt"))
    if sdt is None or not headings:
        return

    holder = next((p for p in sdt.iter(qn("w:p"))
                   if p.find(".//" + qn("w:instrText")) is not None), None)
    if holder is None:
        return

    end_fc = next((fc for fc in holder.iter(qn("w:fldChar"))
                   if fc.get(qn("w:fldCharType")) == "end"), None)
    if end_fc is not None:
        end_fc.getparent().remove(end_fc)

    prev = holder
    for level, title in headings:
        para = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        if level > 1:
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str((level - 1) * 420))
            ppr.append(ind)
        para.append(ppr)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.set(qn("xml:space"), "preserve")
        text.text = title
        run.append(text)
        para.append(run)
        prev.addnext(para)
        prev = para

    tail = OxmlElement("w:p")
    run = OxmlElement("w:r")
    if end_fc is None:
        end_fc = OxmlElement("w:fldChar")
        end_fc.set(qn("w:fldCharType"), "end")
    run.append(end_fc)
    tail.append(run)
    prev.addnext(tail)


def harden_tables(doc) -> None:
    """表头跨页重复、单行不跨页拆分。

    Word 默认允许一行被分页切开，长枚举名（如 DOMAIN_OR_LABEL_MISMATCH）
    会被劈成两半压在下一行。这里给每张表的表头行加 tblHeader、给每一行
    加 cantSplit：跨页时表头跟着走，行整体挪到下一页。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for table in doc.tables:
        for i, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
            if i == 0 and tr_pr.find(qn("w:tblHeader")) is None:
                tr_pr.append(OxmlElement("w:tblHeader"))


def postprocess(path: Path, headings: list[tuple[int, str]]) -> None:
    try:
        import docx
    except ImportError:
        die("找不到 python-docx，请先安装（pip install python-docx）")

    doc = docx.Document(str(path))
    set_fonts(doc)
    set_page_a4(doc)
    move_toc_and_paginate(doc)
    fill_toc(doc, headings)
    harden_tables(doc)
    doc.save(str(path))


def digest(path: Path) -> str:
    """正文 XML 的摘要——zip 时间戳每次都会变，不能直接比整个文件。"""
    with zipfile.ZipFile(path) as z:
        return hashlib.sha256(z.read("word/document.xml")).hexdigest()


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", type=Path, default=DST, help=f"输出路径（默认 {DST}）")
    ap.add_argument("--check", action="store_true", help="只比对，不覆盖现有 docx")
    args = ap.parse_args()

    if not SRC.exists():
        die(f"找不到源文件 {SRC}")

    DST.parent.mkdir(parents=True, exist_ok=True)
    build_md = DOCS / ".report.build.md"
    build_text = strip_manual_toc(SRC.read_text(encoding="utf-8"))
    headings = extract_headings(build_text)
    build_md.write_text(build_text, encoding="utf-8")
    try:
        if args.check:
            target = Path("/tmp/_report_check.docx")
            run_pandoc(build_md, target)
            postprocess(target, headings)
            same = digest(target) == digest(args.out)
            print("一致" if same else "不一致：report.docx 需要重新生成")
            raise SystemExit(0 if same else 1)
        run_pandoc(build_md, args.out)
        postprocess(args.out, headings)
    finally:
        build_md.unlink(missing_ok=True)

    print(f"已生成 {args.out}（源：{SRC.relative_to(ROOT)}，"
          f"{SRC.stat().st_size} 字节 markdown）")


if __name__ == "__main__":
    main()
