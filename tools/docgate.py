#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""DocGate —— 文档生产门禁：把「文稿说了谎」变成红灯 + file:line + 期望/实得。

为什么需要这道门
----------------
2026-10-10 抓到一次真实事故：`docs/figs/` 的配图在撒谎 —— fig2 画着一条代码里
**根本不存在**的 Verify→Repair 回退回路，三张图写「六 Agent」而代码里是 7 个。
**没有任何门禁能发现**，靠人肉看图，5 天后才发现。

根因：图的文字被烘焙进 PNG 像素、报告里的数字全是手写 —— 改代码时没有任何东西
会因此变红。

这道门补的就是这个缺口：文稿里的每一句可机器验证的声明，都在 `docs/facts.json`
里登记；每一次运行都拿文稿 ↔ 清单 ↔ **真实产物** 三方对一遍。

怎么才能让门禁变绿
------------------
只有三条路，且都必须进 diff 让人看见：

  1. **改文稿**，让它说真话（绝大多数情况就是这条）；
  2. **改 facts.json 的 value**，**同时**解释为什么真值变了 —— 但这只在活探针
     （`--live`）也同意时才成立，光改清单是绿不了的；
  3. 那句话确实换了口径，就在 `exempt` 里写清楚它说的不是这件事。

**没有第四条路**：本门永不自动改文稿、永不自动改自己的期望值 ——
否则改不动代码的人会去改期望值，门就白设了。

判据
----
    exit 0    全绿
    exit 1    **真失败**（值不符 / 锚点没了 / 未登记的声明 / 图没重出）
    exit 2    环境不满足（`--strict-env` 时；默认只出声不判死）

「我再也证明不了这句话了」（锚点消失、探针数出 0 个分支）算 **1**，不算 2。
记成 SKIP 会让门禁在重构后安静失效 —— 那正是本次事故的形态。

用法
----
    python3 tools/docgate.py                      # 默认：文稿 ↔ 清单
    python3 tools/docgate.py --live               # 再加：清单 ↔ 真实产物（要语料）
    python3 tools/docgate.py --only facts figures # 只跑某几项
    python3 tools/docgate.py --draft              # 组装期：占位符多了只黄不红
    python3 tools/docgate.py --strict-env         # 环境缺失也算失败（CI 用）
    python3 tools/docgate.py --json /tmp/dg.json  # 机器可读结果
    python3 tools/docgate.py --emit-manifest      # 打印当前锚点/真值，供人 review 后回填

零第三方依赖（只用 stdlib）。
"""

from __future__ import annotations

import argparse
import ast
import datetime as _dt
import json
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import docfacts  # noqa: E402  (同目录的探针模块，tools/ 不是包)

ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "docs" / "facts.json"

CHECKS = ("facts", "mirrors", "figures", "placeholders", "docx", "closure")

FAIL, WARN, SKIP, OK = "fail", "warn", "skip", "ok"


@dataclass
class Finding:
    check: str
    severity: str
    kind: str
    message: str
    file: str | None = None
    line: int | None = None
    expected: object = None
    got: object = None

    def where(self) -> str:
        if not self.file:
            return ""
        return f"{self.file}:{self.line}" if self.line else self.file


@dataclass
class Ctx:
    root: Path
    man: dict
    live: bool = False
    draft: bool = False
    strict_env: bool = False
    json_out: Path | None = None
    _units: dict = field(default_factory=dict)
    _fstr: dict = field(default_factory=dict)

    # -- 文本单元缓存：一个文件只读一次 ------------------------------ #
    def units(self, rel: str):
        if rel not in self._units:
            path = self.root / rel
            if not path.is_file():
                self._units[rel] = None
                self._fstr[rel] = 0
            else:
                u, joined = docfacts.units_for(path, rel)
                self._units[rel] = u
                self._fstr[rel] = joined
        return self._units[rel]

    def fact(self, name: str) -> dict:
        return self.man["facts"][name]

    def sites_of(self, rel: str) -> list[dict]:
        return [s for s in self.man["sites"] if s["file"] == rel]


def _grab(unit_text: str, pattern: str):
    """在一个文本单元上抓数。返回 ``(捕获片段, ...)`` 列表。"""
    rx = re.compile(pattern)
    out = []
    for m in rx.finditer(unit_text):
        if m.groups():
            out.append(m.group(1))
        else:
            out.append(m.group(0))
    return out


def _exempt(text: str, spec: dict) -> str | None:
    for token in spec.get("exempt", ()):
        if token in text:
            return token
    return None


# --------------------------------------------------------------------------- #
# 1) facts：文稿 ↔ 清单 ↔（--live）真实产物
# --------------------------------------------------------------------------- #

# 每个探针收 ctx —— 它探的是**这棵树**，不是"我此刻站在哪"。
# 语料型的探针（index.*）里的路径本就是全局的（data/raw 只有一个），
# 但 AST 型的两条必须认 root，否则 `--root 别的树 --live` 会拿本仓的代码
# 去验另一棵树的清单，给出一句看着很有底气的错话。
PROBES = {
    "index.rules": lambda ctx: docfacts.probe_index()["rules"],
    "index.types": lambda ctx: docfacts.probe_index()["types"],
    "index.attributes": lambda ctx: docfacts.probe_index()["attributes"],
    "index.sources": lambda ctx: docfacts.probe_index()["sources"],
    "guards.py_branches": lambda ctx: docfacts.probe_guards_py_branches(ctx.root),
    "pipeline.agents": lambda ctx: docfacts.probe_pipeline_agents(ctx.root),
}


def check_facts(ctx: Ctx) -> list[Finding]:
    out: list[Finding] = []
    for site in ctx.man["sites"]:
        rel = site["file"]
        fname = site["fact"]
        if fname not in ctx.man["facts"]:
            out.append(Finding("facts", FAIL, "unknown-fact",
                               f"site 引用了不存在的 fact「{fname}」", rel))
            continue
        fact = ctx.fact(fname)
        want = docfacts.normalize_text(fact["value"], fact["norm"])
        units = ctx.units(rel)
        if units is None:
            out.append(Finding("facts", FAIL, "missing-file",
                               f"清单登记的文件不存在（改名/挪走了？）", rel))
            continue

        checked = 0
        bad = False
        for u in units:
            if site["context"] not in u.text:
                continue
            if _exempt(u.text, site):
                continue
            caps = _grab(u.text, site["pattern"])
            if not caps:
                continue          # 这行只是顺嘴提到，不是在声明
            checked += 1
            for cap in caps:
                got = docfacts.normalize(cap, fact["norm"])
                if got != want:
                    bad = True
                    out.append(Finding(
                        "facts", FAIL, "value-mismatch",
                        f"{fname} 与文稿不符（锚点「{site['context']}」，"
                        f"抓到「{cap}」）",
                        rel, u.lineno, want, got))
        if checked == 0:
            out.append(Finding(
                "facts", FAIL, "missing-site",
                f"锚点「{site['context']}」在 {rel} 里找不到任何带数字的单元 —— "
                f"这句话被改写了。跑 `--emit-manifest` 重新登记，"
                f"不是把这里删掉。",
                rel, None, fname, None))
        elif not bad:
            out.append(Finding("facts", OK, "checked",
                               f"{fname} 在 {rel} 的 {checked} 处与清单一致"
                               f"（锚点「{site['context']}」）", rel))

    out.extend(_reverse_scan(ctx))

    if ctx.live:
        out.extend(_live_facts(ctx))
    return out


def _reverse_scan(ctx: Ctx) -> list[Finding]:
    """反着扫一遍：文稿里出现「N 条规则」这类句式的地方，必须已被某个 site 覆盖。

    正向检查只能证明「登记过的话是真的」；这一步才能抓到**新写了一句没人登记的话**。
    """
    out: list[Finding] = []
    spec = ctx.man.get("wait_for", {})
    for g in spec.get("grammars", []):
        rel = g["file"]
        units = ctx.units(rel)
        if units is None:
            continue
        sites_here = ctx.sites_of(rel)
        for u in units:
            if not re.search(g["pattern"], u.text):
                continue
            if _exempt(u.text, g):
                continue
            if any(s["context"] in u.text and re.search(s["pattern"], u.text)
                   for s in sites_here):
                continue          # 已经被某个 site 覆盖
            stmt = _grab(u.text, g["pattern"])
            out.append(Finding(
                "facts", g.get("severity", WARN), "unregistered-claim",
                f"这句话没人登记过：「{u.text.strip()[:70]}」（抓到 {stmt}）—— "
                f"要么在 docs/facts.json 里给它登记一条 site，"
                f"要么在 exempt 里写清它说的不是这件事",
                rel, u.lineno, None, stmt))
    return out


def _live_facts(ctx: Ctx) -> list[Finding]:
    """清单 ↔ 真实产物。没有这一步，光改 facts.json 就能变绿。"""
    out: list[Finding] = []
    for name, fact in ctx.man["facts"].items():
        probe = fact.get("probe")
        if not probe:
            out.append(Finding("facts", WARN, "no-live-probe",
                               f"{name} 是口径型事实，没有活探针可指 —— "
                               f"只验了文稿内部一致，别把它当成验过真值", None))
            continue
        try:
            got = PROBES[probe](ctx)
        except docfacts.ProbeUnavailable as exc:
            out.append(Finding("facts", SKIP, "probe-unavailable",
                               f"{name} 的活探针跑不了：{exc}"))
            continue
        except docfacts.ProbeError as exc:
            out.append(Finding("facts", FAIL, "probe-broken",
                               f"{name} 的活探针坏了：{exc}"))
            continue
        want = docfacts.normalize_text(fact["value"], fact["norm"])
        if got != want:
            out.append(Finding("facts", FAIL, "live-mismatch",
                               f"{name} 与真实产物不符（活探针实测）",
                               None, None, want, got))
    return out


# --------------------------------------------------------------------------- #
# 2) mirrors：两份实现 / 图与代码，必须说同一句话
# --------------------------------------------------------------------------- #

def check_mirrors(ctx: Ctx) -> list[Finding]:
    out: list[Finding] = []
    for spec in ctx.man["mirrors"]:
        if "signatures" in spec:
            out.extend(_mirror_guards(ctx, spec))
        elif "fig" in spec:
            out.extend(_mirror_figure(ctx, spec))
    return out


def _mirror_guards(ctx: Ctx, spec: dict) -> list[Finding]:
    mid = spec["id"]
    out: list[Finding] = []
    py_rel = spec["py"]["file"]
    cpp_rel = spec["cpp"]["file"]
    sigs = spec["signatures"]

    try:
        py_lits = docfacts.func_string_literals(ctx.root, py_rel, spec["py"]["func"])
        py_branches = docfacts.probe_guards_py_branches(ctx.root)
    except docfacts.ProbeError as exc:
        return [Finding("mirrors", FAIL, "probe-broken", f"{mid}: {exc}")]
    try:
        cpp_lits = docfacts.cpp_string_literals(ctx.root / cpp_rel)
        cpp_branches = docfacts.probe_guards_cpp_branches(ctx.root / cpp_rel)
    except docfacts.ProbeUnavailable as exc:
        return [Finding("mirrors", SKIP, "probe-unavailable",
                        f"{mid}: 板端那份查不了（{exc}）—— 跳过 ≠ 一致")]
    except docfacts.ProbeError as exc:
        return [Finding("mirrors", FAIL, "probe-broken", f"{mid}: {exc}")]

    py_text = "\n".join(py_lits)
    cpp_text = "\n".join(cpp_lits)
    for sig in sigs:
        n_py, n_cpp = py_text.count(sig), cpp_text.count(sig)
        if n_py != 1 or n_cpp != 1:
            out.append(Finding(
                "mirrors", FAIL, "signature-drift",
                f"{mid}: 守门签名「{sig}」两侧不是各一次（主机 {n_py} / 板端 {n_cpp}）"
                f"—— 有一侧改了措辞或删了分支",
                None, None, 1, (n_py, n_cpp)))
    if not (py_branches == cpp_branches == len(sigs)):
        out.append(Finding(
            "mirrors", FAIL, "branch-count-drift",
            f"{mid}: 守卫分支数不一致 —— 主机 {py_branches} / 板端 {cpp_branches} / "
            f"登记签名 {len(sigs)}",
            None, None, len(sigs), (py_branches, cpp_branches)))
    if not out:
        out.append(Finding("mirrors", OK, "in-sync",
                           f"{mid}: {py_branches} 道守门在主机与板端逐条对齐"))
    return out


def _mirror_figure(ctx: Ctx, spec: dict) -> list[Finding]:
    mid = spec["id"]
    fig = spec["fig"]
    try:
        want = PROBES[spec["probe"]](ctx)
    except docfacts.ProbeUnavailable as exc:
        return [Finding("mirrors", SKIP, "probe-unavailable", f"{mid}: {exc}")]
    except docfacts.ProbeError as exc:
        return [Finding("mirrors", FAIL, "probe-broken", f"{mid}: {exc}")]

    units = ctx.units(fig["file"])
    if units is None:
        return [Finding("mirrors", FAIL, "missing-file",
                        f"{mid}: {fig['file']} 不在", fig["file"])]
    out, checked = [], 0
    for u in units:
        if fig["context"] not in u.text or _exempt(u.text, fig):
            continue
        caps = _grab(u.text, fig["pattern"])
        if not caps:
            continue
        checked += 1
        for cap in caps:
            got = docfacts.normalize(cap, "cn_or_int")
            if got != want:
                out.append(Finding("mirrors", FAIL, "figure-vs-code",
                                   f"{mid}: 图上写「{cap}」，代码里是 {want}",
                                   fig["file"], u.lineno, want, got))
    if checked == 0:
        out.append(Finding("mirrors", FAIL, "missing-site",
                           f"{mid}: 图脚本里找不到锚点「{fig['context']}」",
                           fig["file"], None, want, None))
    elif not out:
        out.append(Finding("mirrors", OK, "figure-vs-code",
                           f"{mid}: 图上画的 {want} 与代码一致", fig["file"]))
    return out


# --------------------------------------------------------------------------- #
# 3) figures：改了脚本没重新出图
# --------------------------------------------------------------------------- #

def _git(root: Path, *args: str):
    try:
        p = subprocess.run(["git", "-C", str(root), *args],
                           capture_output=True, text=True)
    except OSError:
        return None
    return p.stdout.strip() if p.returncode == 0 else None


def check_figures(ctx: Ctx) -> list[Finding]:
    cfg = ctx.man["figures"]
    script = ctx.root / cfg["script"]
    outs = [ctx.root / o for o in cfg["outputs"]]

    missing = [o for o in cfg["outputs"] if not (ctx.root / o).is_file()]
    if missing:
        return [Finding("figures", FAIL, "missing-output",
                        f"图脚本声明的产物不在：{missing}", cfg["script"])]

    if _git(ctx.root, "rev-parse", "--is-inside-work-tree") != "true":
        return [Finding("figures", SKIP, "no-git",
                        "不是 git 工作树，判不了「改了脚本没出图」")]

    dirty = _git(ctx.root, "status", "--porcelain", "--", cfg["script"])
    if dirty:
        # 工作区里脚本是脏的：git 提交时间还没更新，退化为 mtime。
        newest_png = max(o.stat().st_mtime for o in outs)
        if script.stat().st_mtime > newest_png + 1:
            return [Finding("figures", FAIL, "stale-figure",
                            "改了 make_figs.py 但没重新出图（脚本 mtime 晚于所有 PNG）",
                            cfg["script"])]
        return [Finding("figures", OK, "dirty-but-fresh",
                        "脚本在工作区是脏的，但产物比它新 —— 已重出图，记得提交")]

    t_script = _git(ctx.root, "log", "-1", "--format=%ct", "--", cfg["script"])
    t_pngs = [_git(ctx.root, "log", "-1", "--format=%ct", "--", o)
              for o in cfg["outputs"]]
    if not t_script or any(t is None for t in t_pngs):
        return [Finding("figures", SKIP, "untracked",
                        "脚本或产物还没入库，判不了谱系")]
    newest = max(int(t) for t in t_pngs)
    if int(t_script) > newest:
        when = _dt.datetime.fromtimestamp(int(t_script)).strftime("%Y-%m-%d %H:%M")
        return [Finding("figures", FAIL, "stale-figure",
                        f"make_figs.py 在 {when} 之后改过，但 PNG 没有更新的提交 "
                        f"—— 图还是老的", cfg["script"])]
    return [Finding("figures", OK, "fresh", "图与脚本同一批（或图更新）")]


# --------------------------------------------------------------------------- #
# 4) placeholders：只许变少
# --------------------------------------------------------------------------- #

def check_placeholders(ctx: Ctx) -> list[Finding]:
    out: list[Finding] = []
    today = _dt.date.today()
    for e in ctx.man["placeholders"]["entries"]:
        rel = e["file"]
        path = ctx.root / rel
        if not path.is_file():
            out.append(Finding("placeholders", FAIL, "missing-file",
                               "登记的文件不在", rel))
            continue
        text = path.read_text(encoding="utf-8")
        n = len(re.findall(e["pattern"], text))
        base = e["count"]
        lines = [i for i, ln in enumerate(text.splitlines(), 1)
                 if re.search(e["pattern"], ln)]

        exp = e.get("expires")
        overdue = bool(exp) and n > 0 and _dt.date.fromisoformat(exp) < today

        if n == base and not overdue:
            out.append(Finding("placeholders", OK, "as-expected",
                               f"占位符「{e['pattern']}」= {n} 处，与 baseline 一致"
                               f"（期限 {e.get('expires')}）", rel))
        if n > base:
            sev = WARN if ctx.draft else FAIL
            out.append(Finding("placeholders", sev, "too-many",
                               f"占位符「{e['pattern']}」变多了：{base} → {n} "
                               f"（行 {lines}）。多了就是漏填。",
                               rel, lines[0] if lines else None, base, n))
        elif n < base:
            out.append(Finding("placeholders", WARN, "fewer",
                               f"占位符「{e['pattern']}」{base} → {n} —— "
                               f"填上了就顺手把 facts.json 的 baseline 改成 {n}",
                               rel, None, base, n))

        if overdue:
            sev = WARN if ctx.draft else FAIL
            out.append(Finding("placeholders", sev, "expired",
                               f"占位符「{e['pattern']}」的期限 {exp} 已过，仍剩 {n} 处",
                               rel, None, 0, n))
    return out


# --------------------------------------------------------------------------- #
# 5) docx：派生物必须与源一致
# --------------------------------------------------------------------------- #

def _v3(s: str | None) -> str:
    """版本号归一到前三位。

    ⚠️ 2026-10-23 更正：第四位**会**改变行为 —— 3.1.11 与 3.1.11.1 产出的
    docx `word/document.xml` 不同（实测 --check 直接判不一致）。这里仍只比前三位，
    所以 toolchain-drift 只能挡住大版本；**生成 docx 的本机版本与 CI 钉的版本
    必须逐位相同**（见 .github/workflows/ci.yml 的 docx job 与 facts.json 的 note）。
    """
    if not s:
        return ""
    parts = re.findall(r"\d+", s)[:3]
    return ".".join(parts)


def _pandoc_version() -> str | None:
    exe = shutil.which("pandoc")
    if exe is None:
        return None
    p = subprocess.run([exe, "--version"], capture_output=True, text=True)
    first = p.stdout.splitlines()[0] if p.stdout else ""
    return first.replace("pandoc", "").strip() or None


def _python_docx_version() -> str | None:
    try:
        import docx
    except ImportError:
        return None
    return getattr(docx, "__version__", None) or "unknown"


def check_docx(ctx: Ctx) -> list[Finding]:
    cfg = ctx.man["docx"]
    artifact = ctx.root / cfg["artifact"]
    if not artifact.is_file():
        return [Finding("docx", SKIP, "missing-artifact",
                        f"{cfg['artifact']} 还没生成过")]
    if shutil.which("pandoc") is None:
        return [Finding("docx", SKIP, "no-pandoc", "本机没装 pandoc")]
    got_docx = _python_docx_version()
    if got_docx is None:
        return [Finding("docx", SKIP, "no-python-docx", "本机没装 python-docx")]

    want = cfg["provenance_env"]
    got_pandoc = _pandoc_version()
    if (_v3(got_pandoc) != _v3(want.get("pandoc"))
            or _v3(got_docx) != _v3(want.get("python_docx"))):
        return [Finding("docx", SKIP, "toolchain-drift",
                        f"工具链与登记不符 —— 登记 pandoc {want.get('pandoc')} / "
                        f"python-docx {want.get('python_docx')}，"
                        f"本机 pandoc {got_pandoc} / python-docx {got_docx}。"
                        f"版本会改变正文 XML 细节，这是环境差异不是文档缺陷。")]

    builder = ctx.root / cfg["builder"]
    if not builder.is_file():
        return [Finding("docx", FAIL, "missing-builder",
                        f"生成脚本不在：{cfg['builder']}")]
    p = subprocess.run([sys.executable, str(builder), "--check"],
                       cwd=str(ctx.root), capture_output=True, text=True)
    msg = (p.stdout + p.stderr).strip().splitlines()
    tail = msg[-1] if msg else f"退出码 {p.returncode}"
    if p.returncode == 0:
        return [Finding("docx", OK, "consistent", "report.docx 与 report.md 一致")]
    if p.returncode == 1:
        return [Finding("docx", FAIL, "inconsistent",
                        f"{cfg['artifact']} 落后于 {cfg['md']}（{tail}）—— "
                        f"跑 `python3 {cfg['builder']}` 重新生成")]
    return [Finding("docx", FAIL, "builder-broken",
                    f"生成脚本以退出码 {p.returncode} 结束：{tail}")]


# --------------------------------------------------------------------------- #
# 6) closure：门禁自证 —— 别让它自己被掏空
# --------------------------------------------------------------------------- #

def check_closure(ctx: Ctx, seen: set[str] | None = None) -> list[Finding]:
    out: list[Finding] = []

    for rel in ctx.man["managed"]:
        if not (ctx.root / rel).is_file():
            out.append(Finding("closure", FAIL, "missing-managed",
                               "受管文件不在", rel))
        else:
            ctx.units(rel)          # 顺带把 f-string 计数读全

    if not ctx.man["facts"]:
        out.append(Finding("closure", FAIL, "empty-facts",
                           "facts 是空的 —— 门禁没有可验的东西了"))

    facts_with_site = {s["fact"] for s in ctx.man["sites"]}
    for name, fact in ctx.man["facts"].items():
        if name not in facts_with_site:
            out.append(Finding("closure", FAIL, "fact-without-site",
                               f"fact「{name}」没有任何 site —— 它从未被验过"))
        if fact.get("probe") and fact["probe"] not in PROBES:
            out.append(Finding("closure", FAIL, "unknown-probe",
                               f"fact「{name}」的活探针「{fact['probe']}」不存在"))
    if not ctx.man["sites"]:
        out.append(Finding("closure", FAIL, "empty-sites", "sites 是空的"))

    # 报告里引用的图必须存在；figs/ 里也不该有没人引用的孤儿 PNG
    cfg = ctx.man["figures"]
    declared = {Path(o).name for o in cfg["outputs"]}
    referenced: set[str] = set()
    for rel in cfg["consumers"]:
        units = ctx.units(rel) or []
        for u in units:
            for m in re.finditer(r"!\[[^\]]*\]\(figs/([^)]+)\)", u.text):
                referenced.add(m.group(1))
                if not (ctx.root / "docs" / "figs" / m.group(1)).is_file():
                    out.append(Finding("closure", FAIL, "missing-figure",
                                       f"报告引用了不存在的图：{m.group(1)}",
                                       rel, u.lineno))
    orphans = declared - referenced
    if orphans:
        out.append(Finding("closure", WARN, "orphan-figure",
                           f"生成了却没人引用的图：{sorted(orphans)}"))
    unused = referenced - declared
    if unused:
        out.append(Finding("closure", WARN, "undeclared-figure",
                           f"报告引用了但图脚本没声明的图：{sorted(unused)}"))

    # f-string：静态片段是按单元逐个比对的，跨片段的句子会走 missing-site，
    # 所以这里不是风险点，只是一个「比到了多少」的计数，出声以免被误当漏检。
    total_fstr = sum(n for n in ctx._fstr.values())
    if total_fstr:
        out.append(Finding("closure", OK, "fstring-note",
                           f"受管文件里共 {total_fstr} 处 f-string；"
                           f"其静态片段逐个当单元比对，跨片段的句子由 missing-site 兜住"))

    if seen is not None:
        for name in CHECKS:
            if name not in seen:
                out.append(Finding("closure", FAIL, "check-not-run",
                                   f"check「{name}」这一轮没产出结果"))
    return out


# --------------------------------------------------------------------------- #
# --emit-manifest：唯一刷新通道，且只写 stdout
# --------------------------------------------------------------------------- #

def emit_manifest(ctx: Ctx) -> int:
    print(json.dumps({
        "updated": _dt.date.today().isoformat(),
        "facts": {n: {"value": f["value"]}
                  for n, f in ctx.man["facts"].items()},
        "sites": [
            {"fact": s["fact"], "file": s["file"], "context": s["context"],
             "pattern": s["pattern"],
             "hits": _site_hits(ctx, s)}
            for s in ctx.man["sites"]
        ],
    }, ensure_ascii=False, indent=2))
    print("\n（以上只是现状快照。要改 facts.json，请人工 review 后手动编辑；"
          "本门不会替你写盘。）", file=sys.stderr)
    return 0


def _site_hits(ctx: Ctx, site: dict) -> list[dict]:
    units = ctx.units(site["file"]) or []
    hits = []
    for u in units:
        if site["context"] not in u.text:
            continue
        caps = _grab(u.text, site["pattern"])
        if caps:
            hits.append({"line": u.lineno, "grabbed": caps})
    return hits


# --------------------------------------------------------------------------- #
# 跑
# --------------------------------------------------------------------------- #

def load_manifest(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def run(ctx: Ctx, only: list[str] | None = None) -> dict:
    order = [c for c in CHECKS if not only or c in only]
    results: dict[str, list[Finding]] = {}
    for name in order:
        if name == "closure":
            continue
        results[name] = {
            "facts": check_facts, "mirrors": check_mirrors,
            "figures": check_figures, "placeholders": check_placeholders,
            "docx": check_docx,
        }[name](ctx)
    if "closure" in order:
        results["closure"] = check_closure(ctx, set(results) | {"closure"})
    return results


def worst(results: dict) -> str:
    sev = {OK: 0, SKIP: 1, WARN: 2, FAIL: 3}
    rank = max((sev[f.severity] for fs in results.values() for f in fs),
               default=0)
    return {0: OK, 1: SKIP, 2: WARN, 3: FAIL}[rank]


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--root", type=Path, default=ROOT)
    ap.add_argument("--manifest", type=Path, default=None)
    ap.add_argument("--live", action="store_true",
                    help="再验一层：清单 ↔ 真实产物（需要 data/raw 语料）")
    ap.add_argument("--draft", action="store_true",
                    help="组装期：占位符多了只警告，不判失败")
    ap.add_argument("--strict-env", action="store_true",
                    help="环境不满足也算失败（exit 2）")
    ap.add_argument("--only", nargs="+", choices=list(CHECKS), default=None)
    ap.add_argument("--json", dest="json_out", type=Path, default=None)
    ap.add_argument("--emit-manifest", action="store_true")
    args = ap.parse_args(argv)

    man_path = args.manifest or (args.root / "docs" / "facts.json")
    if not man_path.is_file():
        print(f"错误：找不到真值清单 {man_path}", file=sys.stderr)
        return 1
    ctx = Ctx(root=args.root, man=load_manifest(man_path),
              live=args.live, draft=args.draft, strict_env=args.strict_env,
              json_out=args.json_out)

    if args.emit_manifest:
        return emit_manifest(ctx)

    results = run(ctx, args.only)

    icon = {OK: " ok ", FAIL: "FAIL", WARN: "warn", SKIP: "SKIP"}
    n_fail = n_warn = n_skip = 0
    for name, fs in results.items():
        bad = [f for f in fs if f.severity != OK]
        head = f"{len(fs) - len(bad)}/{len(fs)} 通过"
        print(f"\n=== {name}  ({head}) ===")
        if not fs:
            print("  (无登记项)")
        for f in fs:
            if f.severity == OK:
                print(f"  [ ok ] {f.message}")
                continue
            tag = f"[{icon[f.severity]}]"
            if f.expected is not None and f.got is not None:
                gap = f"  期望 {f.expected!r} / 实得 {f.got!r}"
            elif f.expected is not None:
                gap = f"  期望 {f.expected!r}"
            elif f.got is not None:
                gap = f"  实得 {f.got!r}"
            else:
                gap = ""
            loc = f"  @ {f.where()}" if f.where() else ""
            print(f"  {tag} {f.kind}: {f.message}{gap}{loc}")
            n_fail += f.severity == FAIL
            n_warn += f.severity == WARN
            n_skip += f.severity == SKIP

    verdict = worst(results)
    age = man_path and ctx.man.get("updated")
    print(f"\n合计：{n_fail} 失败 / {n_warn} 警告 / {n_skip} 跳过"
          f"   →  {verdict.upper()}")
    if n_skip:
        print("注意：跳过不算通过 —— 上面每一项跳过都写在明处了。")
    if age:
        days = (_dt.date.today() - _dt.date.fromisoformat(age)).days
        print(f"真值清单 updated={age}（{days} 天前）。")

    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(json.dumps({
            "verdict": verdict,
            "counts": {"fail": n_fail, "warn": n_warn, "skip": n_skip},
            "manifest_updated": age,
            "checks": {k: [f.__dict__ for f in v] for k, v in results.items()},
        }, ensure_ascii=False, indent=2), encoding="utf-8")

    if n_fail:
        return 1
    if n_skip and args.strict_env:
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
